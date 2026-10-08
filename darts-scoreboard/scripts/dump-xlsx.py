# Dump an Excel sheet as a coloured text grid: value + fill colour per cell.
# Usage: python scripts/dump-xlsx.py "file.xlsx" "SheetName" [max_row] [max_col]
import sys
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

path, sheet = sys.argv[1], sys.argv[2]
max_row = int(sys.argv[3]) if len(sys.argv) > 3 else 45
max_col = int(sys.argv[4]) if len(sys.argv) > 4 else 22

wb = load_workbook(path, data_only=True)
ws = wb[sheet]

def rgb(cell):
    f = cell.fill
    if f is None or f.fgColor is None or f.patternType is None:
        return '...'
    c = f.fgColor
    if c.type == 'rgb' and isinstance(c.rgb, str):
        return c.rgb[-6:].upper()
    if c.type == 'indexed' and isinstance(c.indexed, int):
        return f'idx{c.indexed}'
    if c.type == 'theme':
        return f'thm{c.theme}'
    return '...'

for r in range(1, min(ws.max_row, max_row) + 1):
    line = []
    for c in range(1, min(ws.max_column, max_col) + 1):
        cell = ws.cell(row=r, column=c)
        v = cell.value
        v = '' if v is None else str(v)[:10]
        line.append(f'{v}|{rgb(cell)}')
    if any(l.split('|')[0] for l in line) or any(l.split('|')[1] != '...' for l in line):
        print(f'{r:>3} ' + ' '.join(f'{l:<14}' for l in line))
